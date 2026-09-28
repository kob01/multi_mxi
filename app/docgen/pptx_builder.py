"""PPT(.pptx) 构建器: python-pptx, 纯同步函数(由 tools/docgen.py 用 to_thread 卸载)。

spec 结构:
    {"title": "封面标题", "subtitle": "封面副标题",
     "slides": [{"title": "页标题", "bullets": ["要点1", {..可再带 "level": 1 缩进}]}],
     "images": [{"path": "本地png", "caption": "图1", "width": 6.5}]}

images 由调用方预先解析为本地 PNG 绝对路径; 每张图单独占一页(空白版式), 居中 + 下方题注。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def _bullet_text(item: Any) -> tuple[str, int]:
    if isinstance(item, dict):
        return str(item.get("text") or ""), max(0, min(int(item.get("level") or 0), 4))
    return str(item or ""), 0


def build_pptx(path: Path, spec: dict[str, Any]) -> None:
    from pptx import Presentation

    slides_spec = [s for s in (spec.get("slides") or []) if isinstance(s, dict)]
    title = str(spec.get("title") or "").strip()
    if not title and not slides_spec:
        raise ValueError("spec 至少需要 title 或 slides, 否则生成的是空演示文稿")

    prs = Presentation()
    # 封面: 版式 0(TITLE) 自带标题+副标题占位符; 没给 title 时不做封面页。
    if title:
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = title
        subtitle = str(spec.get("subtitle") or "").strip()
        placeholders = [ph for ph in cover.placeholders if ph.placeholder_format.idx == 1]
        if subtitle and placeholders:
            placeholders[0].text = subtitle

    for slide_spec in slides_spec:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # TITLE_AND_CONTENT
        slide.shapes.title.text = str(slide_spec.get("title") or "").strip()
        body = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)
        if body is None:
            continue
        frame = body.text_frame
        first = True
        for item in slide_spec.get("bullets") or []:
            text, level = _bullet_text(item)
            if not text.strip():
                continue
            para = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            para.text = text.strip()
            para.level = level

    _add_images(prs, spec.get("images") or [])

    prs.save(str(path))


def _add_images(prs: Any, images: list[dict[str, Any]]) -> None:
    """每张图新开一页(空白版式 6), 图居中, 下方一个题注文本框。

    默认幻灯片按 10x7.5 英寸算; width 英寸缺省 6.5, 限在 1~9。图不可读则跳过这张。
    """
    from pptx.util import Inches

    for img in images:
        p = str((img or {}).get("path") or "")
        if not p or not Path(p).is_file():
            continue
        try:
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        except Exception:  # noqa: BLE001 - 版式缺失不致命, 放弃这一张
            continue
        try:
            width = float(img.get("width") or 6.5)
        except (TypeError, ValueError):
            width = 6.5
        width = min(max(width, 1.0), 9.0)
        left = Inches(max(0.0, (10.0 - width) / 2))
        try:
            pic = slide.shapes.add_picture(p, left, Inches(1.0), width=Inches(width))
            top_after = (pic.top + pic.height) / 914400 if pic.height else 3.5
        except Exception:  # noqa: BLE001 - 坏图不毁整个演示文稿
            continue
        caption = str(img.get("caption") or "").strip()
        if caption:
            box = slide.shapes.add_textbox(left, Inches(min(top_after + 0.2, 6.6)), Inches(width), Inches(0.6))
            box.text_frame.text = caption
