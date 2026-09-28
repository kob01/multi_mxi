"""Word(.docx) 构建器: python-docx, 纯同步函数(由 tools/docgen.py 用 to_thread 卸载)。

spec 结构(全部可选, 缺省即跳过):
    {"title": "标题", "subtitle": "副标题",
     "sections": [{"heading": "小节标题", "body": "正文, 换行分段"}],
     "bullets": ["要点1", "要点2"],
     "table": {"columns": ["列1", "列2"], "rows": [["a", "b"]]},
     "images": [{"path": "本地png", "caption": "图1", "width": 5.5}]}

images 由调用方(app/tools/docgen.py 经 app/docgen/images.py)预先解析为本地 PNG 绝对路径,
builder 只按路径嵌入 —— 网络/SSRF/归一化都不在这一层。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def build_docx(path: Path, spec: dict[str, Any]) -> None:
    from docx import Document

    if not str(spec.get("title") or "").strip() and not spec.get("sections"):
        raise ValueError("spec 至少需要 title 或 sections, 否则生成的是空文档")

    doc = Document()
    title = str(spec.get("title") or "").strip()
    if title:
        doc.add_heading(title, level=0)
    subtitle = str(spec.get("subtitle") or "").strip()
    if subtitle:
        para = doc.add_paragraph()
        run = para.add_run(subtitle)
        run.italic = True

    for section in spec.get("sections") or []:
        heading = str((section or {}).get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for line in str((section or {}).get("body") or "").splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())

    for bullet in spec.get("bullets") or []:
        text = str(bullet).strip()
        if not text:
            continue
        try:
            doc.add_paragraph(text, style="List Bullet")
        except KeyError:
            # 个别模板缺该样式时退化为带圆点前缀的普通段落, 不让样式名毁掉整份文档。
            doc.add_paragraph(f"• {text}")

    table = spec.get("table") or {}
    columns = [str(c) for c in (table.get("columns") or [])]
    if columns:
        rows = [[str(c) for c in (row or [])] for row in (table.get("rows") or [])]
        grid = doc.add_table(rows=1, cols=len(columns))
        try:
            grid.style = "Table Grid"
        except KeyError:
            pass
        for idx, name in enumerate(columns):
            grid.rows[0].cells[idx].text = name
        for row in rows:
            cells = grid.add_row().cells
            for idx in range(len(columns)):
                cells[idx].text = row[idx] if idx < len(row) else ""

    _add_images(doc, spec.get("images") or [])

    doc.save(str(path))


def _add_images(doc: Any, images: list[dict[str, Any]]) -> None:
    """按解析好的本地 PNG 路径逐张嵌入(图后跟一行斜体题注)。

    width 单位是英寸(缺省 5.5, 约 A4 正文宽); 图不可读(文件丢失/格式不支持)时
    跳过这一张而不是抛异常毁掉整份文档 —— 与图片解析层"能降就降"一致。
    """
    from docx.shared import Inches

    for img in images:
        p = str((img or {}).get("path") or "")
        if not p or not Path(p).is_file():
            continue
        try:
            width = float(img.get("width") or 5.5)
        except (TypeError, ValueError):
            width = 5.5
        try:
            doc.add_picture(p, width=Inches(min(max(width, 1.0), 6.5)))
        except Exception:  # noqa: BLE001 - 单张图坏不影响其余内容与整份文档
            continue
        caption = str(img.get("caption") or "").strip()
        if caption:
            para = doc.add_paragraph()
            run = para.add_run(caption)
            run.italic = True
